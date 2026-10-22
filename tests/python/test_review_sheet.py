"""tools/review_sheet.py export -> (reviewer edits with openpyxl) -> import, in a temp project."""
import datetime
import json
import os
import shutil
import time
import unittest

from openpyxl import load_workbook

from . import helpers as h
from tpcloc import translations

SRC, OTHER = "I2LS_Test", "I2LS_Other"
ROWS = {
    SRC: [
        ("UI/Mods/A", "Mod A text", ""),
        ("UI/Mods/B", "Mod B text", ""),
        ("UI/Mods/C", "Mod C text", ""),
        ("UI/Mods/D", "Mod D text", ""),  # no translation yet
        ("X/1", "First\nSecond", ""),
        ("X/2", "Edited one", ""),
        ("X/3", "Third one", ""),
    ],
    OTHER: [("O/Ghost", "Ghost thing", "")],
}
CATS = {(SRC, k): "Mods" for k in ("UI/Mods/A", "UI/Mods/B", "UI/Mods/C", "UI/Mods/D")}
CATS[(OTHER, "O/Ghost")] = "DLC_Ghost"


def ours():
    return {
        SRC: {
            "UI/Mods/A": h.entry("MOD A訳"),
            "UI/Mods/B": h.entry("MOD B訳", note="訳注B"),
            "UI/Mods/C": h.entry("MOD C訳", "reviewed", review="メモC"),
            "X/1": h.entry("  一行目\n二行目\n", note="前後の空白"),
            "X/2": h.entry("修正版", "edited"),
            "X/3": h.entry("三つ目"),
        },
        OTHER: {"O/Ghost": h.entry("ゴーストのもの")},
    }


IDS = [f"{SRC}|UI/Mods/A", f"{SRC}|UI/Mods/B", f"{SRC}|UI/Mods/C", f"{SRC}|UI/Mods/D", f"{SRC}|X/1",
       f"{SRC}|X/2", f"{SRC}|X/3", f"{OTHER}|O/Ghost"]
TRANSLATED = len(IDS) - 1


class Sheet:
    """Reviewer-side editing of an exported workbook (openpyxl, like saving from Excel)."""

    def __init__(self, path):
        self.path = path
        self.wb = load_workbook(path)

    def find(self, ref, nth=0):
        hits = []
        for ws in self.wb.worksheets:
            if ws.cell(1, 1).value != "ID":
                continue
            for r in range(2, ws.max_row + 1):
                if ws.cell(r, 1).value == ref:
                    hits.append((ws, r))
        return hits[nth]

    def col(self, ws, name):
        for c in range(1, ws.max_column + 1):
            if ws.cell(1, c).value == name:
                return c
        raise KeyError(name)

    def get(self, ref, name):
        ws, r = self.find(ref)
        return ws.cell(r, self.col(ws, name)).value

    def set(self, ref, name, value, nth=0):
        ws, r = self.find(ref, nth)
        ws.cell(r, self.col(ws, name)).value = value

    def save(self):
        self.wb.save(self.path)


class ReviewBase(unittest.TestCase):
    def setUp(self):
        self.root = h.make_project(h.new_dir(self), ROWS, ours=ours(), categories=CATS)
        self.xlsx = self.root / "review" / "sheet.xlsx"

    def export(self, *args):
        return h.run_tool("review_sheet.py", "export", "--out", self.xlsx, *args, project=self.root)

    def do_import(self, *args):
        r = h.run_tool("review_sheet.py", "import", self.xlsx, *args, project=self.root)
        self.assertEqual(r.code, 0, r)
        return r

    def json(self, source=SRC):
        return h.load_ours(self.root)[source]

    def save_json(self, source, entries):
        with h.project_paths(self.root):
            translations.save(source, entries, [k for k, _, _ in ROWS[source]])

    def stats(self, r):
        line = next(l for l in r.out.splitlines() if l.startswith("修正 "))
        return line

    def exported(self):
        r = self.export()
        self.assertEqual(r.code, 0, r)
        return Sheet(self.xlsx)


class Export(ReviewBase):
    def test_layout(self):
        r = self.export()
        self.assertEqual(r.code, 0, r)
        self.assertIn(f"{len(IDS)} rows -> {self.xlsx}", r.out)
        self.assertIn("{'MOD画面': 4, 'スクールスピリッツ': 1, 'その他': 3}", r.out)
        wb = load_workbook(self.xlsx)
        self.assertEqual(wb.sheetnames, ["使い方", "MOD画面", "スクールスピリッツ", "その他"])
        ws = wb["MOD画面"]
        header = [c.value for c in ws[1]]
        self.assertEqual(header, ["ID", "ファイル", "種類", "英語（原文）", "日本語訳", "状態", "訳注", "レビューメモ",
                                  "書き出し時ID", "書き出し時の日本語訳", "書き出し時の状態", "書き出し時の訳注",
                                  "書き出し時のレビューメモ"])
        for letter in "IJKLM":
            self.assertTrue(ws.column_dimensions[letter].hidden, letter)
        rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
        self.assertEqual(rows[2], [f"{SRC}|UI/Mods/C", "Test", "ラベル", "Mod C text", "MOD C訳", "確認済み", None,
                                   "メモC", f"{SRC}|UI/Mods/C", "MOD C訳", "確認済み", None, "メモC"])
        self.assertEqual(rows[3][4:8], [None, "未確認", None, None])  # untranslated row
        other = [[c.value for c in row] for row in wb["その他"].iter_rows(min_row=2)]
        self.assertEqual(other[0][4], "  一行目\n二行目\n")
        self.assertEqual(other[1][5], "修正済み")
        self.assertEqual(len(ws.data_validations.dataValidation), 1)
        self.assertEqual(ws.freeze_panes, "D2")
        self.assertTrue(any(isinstance(c.value, str) and c.value.startswith("作成日時: ") for c in wb["使い方"]["A"]))

    def test_refuses_to_overwrite_without_force(self):
        self.export()
        before = self.xlsx.read_bytes()
        r = self.export()
        self.assertEqual(r.code, 1, r)
        self.assertIn("--force", r.out)
        self.assertEqual(self.xlsx.read_bytes(), before)
        time.sleep(1.1)  # the zip timestamps change
        r = self.export("--force")
        self.assertEqual(r.code, 0, r)
        self.assertNotEqual(self.xlsx.read_bytes(), before)

    def test_export_does_not_touch_translations(self):
        before = h.snapshot_files(self.root / "translations")
        self.export()
        self.assertEqual(h.snapshot_files(self.root / "translations"), before)


class Import(ReviewBase):
    def test_no_op_round_trip(self):
        self.exported()
        before = h.snapshot_files(self.root / "translations")
        mtimes = {p: p.stat().st_mtime_ns for p in (self.root / "translations" / "ja").iterdir()}
        r = self.do_import()
        self.assertEqual(self.stats(r), f"修正 0 件、確認済み 0 件、メモ 0 件、変更なし {TRANSLATED} 件、競合 0 件")
        self.assertEqual(r.lines("PROBLEM"), [f"PROBLEM {SRC}|UI/Mods/D: 翻訳データにありません"])
        self.assertEqual(r.lines("CONFLICT"), [])
        self.assertEqual(h.snapshot_files(self.root / "translations"), before)
        self.assertEqual({p: p.stat().st_mtime_ns for p in mtimes}, mtimes)  # not even rewritten

    def test_no_op_after_resave(self):
        s = self.exported()
        s.save()  # openpyxl round trip of the whole workbook (as a reviewer's save would do)
        before = h.snapshot_files(self.root / "translations")
        r = self.do_import()
        self.assertIn("競合 0 件", self.stats(r))
        self.assertEqual(h.snapshot_files(self.root / "translations"), before)

    def test_reviewer_edits(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "MOD A 直した訳\n改行")
        s.set(f"{SRC}|UI/Mods/B", "状態", " 確認済み ")
        s.set(f"{SRC}|UI/Mods/C", "レビューメモ", None)  # memo cleared
        s.set(f"{SRC}|X/3", "レビューメモ", "ここが気になる")
        s.set(f"{SRC}|X/3", "訳注", "新しい訳注")
        s.set(f"{SRC}|X/2", "状態", "未確認")
        s.set(f"{OTHER}|O/Ghost", "状態", "修正済み")
        s.save()
        r = self.do_import()
        self.assertEqual(self.stats(r), "修正 1 件、確認済み 1 件、メモ 1 件、変更なし 1 件、競合 0 件")
        j = self.json()
        self.assertEqual(j["UI/Mods/A"], {"ja": "MOD A 直した訳\n改行", "status": "edited", "note": ""})
        self.assertEqual(j["UI/Mods/B"], h.entry("MOD B訳", "reviewed", "訳注B"))
        self.assertEqual(j["UI/Mods/C"], h.entry("MOD C訳", "reviewed"))
        self.assertEqual(j["X/3"], h.entry("三つ目", note="新しい訳注", review="ここが気になる"))
        self.assertEqual(j["X/2"], h.entry("修正版", "draft"))
        self.assertEqual(j["X/1"], ours()[SRC]["X/1"])
        self.assertEqual(self.json(OTHER)["O/Ghost"], h.entry("ゴーストのもの", "edited"))
        self.assertEqual(list(j), list(ours()[SRC]))  # game key order kept

    def test_edit_and_state_in_one_row(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "新訳")
        s.set(f"{SRC}|UI/Mods/A", "状態", "確認済み")
        s.save()
        self.do_import()
        self.assertEqual(self.json()["UI/Mods/A"]["status"], "edited")

    def test_json_changed_after_export_is_kept_when_sheet_untouched(self):
        self.exported()
        j = self.json()
        j["UI/Mods/A"] = h.entry("JSONで直した", "edited", "JSONの訳注", review="JSONのメモ")
        self.save_json(SRC, j)
        r = self.do_import()
        self.assertEqual(r.lines("CONFLICT"), [])
        self.assertEqual(self.json()["UI/Mods/A"], j["UI/Mods/A"])

    def test_conflicts(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "シートの訳")
        s.set(f"{SRC}|UI/Mods/B", "状態", "確認済み")
        s.set(f"{SRC}|X/3", "訳注", "シートの訳注")
        s.set(f"{SRC}|X/2", "レビューメモ", "シートのメモ")
        s.set(f"{SRC}|X/1", "状態", "確認済み")
        s.save()
        j = self.json()
        j["UI/Mods/A"]["ja"] = "JSONの訳"
        j["UI/Mods/B"]["status"] = "edited"
        j["X/3"]["note"] = "JSONの訳注"
        j["X/2"]["review"] = "JSONのメモ"
        j["X/1"]["ja"] = "JSONで直した"
        self.save_json(SRC, j)
        before = h.snapshot_files(self.root / "translations")
        r = self.do_import()
        self.assertIn("競合 5 件", self.stats(r))
        conflicts = r.lines("CONFLICT I2LS_")
        self.assertEqual(len(conflicts), 5, r)
        self.assertTrue(any("UI/Mods/A: 日本語訳がシートと翻訳データの両方で変わっています（シート「シートの訳」／書き出し時「MOD A訳」／現在「JSONの訳」）" in c
                            for c in conflicts), conflicts)
        self.assertTrue(any("UI/Mods/B: 状態がシートと翻訳データの両方で変わっています" in c for c in conflicts))
        self.assertTrue(any("X/3: 訳注がシートと翻訳データの両方で変わっています" in c for c in conflicts))
        self.assertTrue(any("X/2: レビューメモがシートと翻訳データの両方で変わっています" in c for c in conflicts))
        self.assertTrue(any("X/1: 状態を「確認済み」にした訳が、書き出し後に翻訳データで変わっています" in c for c in conflicts))
        self.assertIn("CONFLICT の項目は取り込んでいません", r.out)
        self.assertEqual(h.snapshot_files(self.root / "translations"), before)

    def test_same_change_on_both_sides_is_not_a_conflict(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "同じ訳")
        s.save()
        j = self.json()
        j["UI/Mods/A"] = h.entry("同じ訳", "edited")
        self.save_json(SRC, j)
        r = self.do_import()
        self.assertIn("競合 0 件", self.stats(r))
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("同じ訳", "edited"))

    def test_conflicting_row_is_left_alone(self):
        """A row reported as CONFLICT ("CONFLICT の項目は取り込んでいません") is not changed at all (RS1 fixed)."""
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "シートの訳")
        s.set(f"{SRC}|UI/Mods/A", "訳注", "シートの訳注")
        s.set(f"{SRC}|UI/Mods/A", "レビューメモ", "シートのメモ")
        s.save()
        j = self.json()
        j["UI/Mods/A"]["ja"] = "JSONの訳"
        self.save_json(SRC, j)
        before = h.snapshot_files(self.root / "translations")
        r = self.do_import()
        # the note/memo of the conflicting row are neither applied nor counted
        self.assertEqual(self.stats(r), f"修正 0 件、確認済み 0 件、メモ 0 件、変更なし {TRANSLATED - 1} 件、競合 1 件")
        self.assertEqual(len(r.lines("CONFLICT I2LS_")), 1, r)
        self.assertIn("CONFLICT の項目は取り込んでいません", r.out)
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("JSONの訳"))
        self.assertEqual(h.snapshot_files(self.root / "translations"), before)  # not even rewritten

    def test_conflict_in_one_field_blocks_the_other_fields_of_the_row(self):
        """Only the memo conflicts; the edited translation and state of that row are not imported either."""
        s = self.exported()
        s.set(f"{SRC}|X/3", "日本語訳", "シートで直した三つ目")
        s.set(f"{SRC}|X/3", "レビューメモ", "シートのメモ")
        s.set(f"{SRC}|UI/Mods/B", "日本語訳", "競合しない行の修正")
        s.save()
        j = self.json()
        j["X/3"]["review"] = "JSONのメモ"
        self.save_json(SRC, j)
        r = self.do_import()
        self.assertEqual(self.stats(r), f"修正 1 件、確認済み 0 件、メモ 0 件、変更なし {TRANSLATED - 2} 件、競合 1 件")
        conflicts = r.lines("CONFLICT I2LS_")
        self.assertEqual(len(conflicts), 1, r)
        self.assertIn("X/3: レビューメモがシートと翻訳データの両方で変わっています", conflicts[0])
        got = self.json()
        self.assertEqual(got["X/3"], h.entry("三つ目", review="JSONのメモ"))
        self.assertEqual(got["UI/Mods/B"], h.entry("競合しない行の修正", "edited", "訳注B"))

    def test_conflicting_row_in_a_source_that_has_other_changes(self):
        """The source file is rewritten for the other row, but the conflicting row keeps its JSON values."""
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "シートの訳")
        s.set(f"{SRC}|UI/Mods/A", "状態", "確認済み")
        s.set(f"{SRC}|UI/Mods/C", "レビューメモ", "新しいメモC")
        s.save()
        j = self.json()
        j["UI/Mods/A"] = h.entry("JSONの訳", "edited", "JSONの訳注")
        self.save_json(SRC, j)
        r = self.do_import()
        self.assertIn("メモ 1 件", self.stats(r))
        self.assertIn("競合 1 件", self.stats(r))
        got = self.json()
        self.assertEqual(got["UI/Mods/A"], h.entry("JSONの訳", "edited", "JSONの訳注"))
        self.assertEqual(got["UI/Mods/C"], h.entry("MOD C訳", "reviewed", review="新しいメモC"))
        self.assertEqual(list(got), list(ours()[SRC]))

    def test_state_blank_or_invalid(self):
        for label in (None, "OK", "確認"):
            with self.subTest(label=label):
                self.save_json(SRC, ours()[SRC])
                self.assertEqual(self.export("--force").code, 0)
                s = Sheet(self.xlsx)
                s.set(f"{SRC}|UI/Mods/B", "状態", label)
                s.set(f"{SRC}|UI/Mods/B", "レビューメモ", f"メモ{label}")
                s.save()
                r = self.do_import()
                want = "" if label is None else label
                self.assertIn(f"PROBLEM {SRC}|UI/Mods/B: 状態「{want}」は未確認・確認済み・修正済みのどれでもないので、状態は変えませんでした",
                              r.out)
                e = self.json()["UI/Mods/B"]
                self.assertEqual(e["status"], "draft")
                self.assertEqual(e["review"], f"メモ{label}")  # the rest of the row is still imported

    def test_empty_translation_refused(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "   ")
        s.save()
        r = self.do_import()
        self.assertIn(f"PROBLEM {SRC}|UI/Mods/A: 日本語訳が空なので取り込みませんでした", r.out)
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("MOD A訳"))

    def test_formula_number_and_date_cells(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "=E3")
        s.set(f"{SRC}|UI/Mods/B", "訳注", 123)
        s.set(f"{SRC}|X/3", "レビューメモ", datetime.datetime(2026, 1, 2))
        s.set(f"{SRC}|X/2", "状態", True)
        s.save()
        before = h.snapshot_files(self.root / "translations")
        r = self.do_import()
        problems = r.lines("PROBLEM")
        for ref, what in ((f"{SRC}|UI/Mods/A", "日本語訳"), (f"{SRC}|UI/Mods/B", "訳注"), (f"{SRC}|X/3", "レビューメモ"),
                          (f"{SRC}|X/2", "状態")):
            self.assertIn(f"PROBLEM {ref}: {what}が数式・数値・日付などになっているので取り込みませんでした（文字で入力し直してください）",
                          problems)
        self.assertEqual(h.snapshot_files(self.root / "translations"), before)

    def test_formula_in_hidden_column(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "書き出し時の日本語訳", "=E2")
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "新訳")
        s.save()
        r = self.do_import()
        self.assertIn(f"PROBLEM {SRC}|UI/Mods/A: 書き出し時の値（非表示の列）が無いか食い違っているので取り込みませんでした", r.out)
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("MOD A訳"))

    def test_duplicate_ids(self):
        s = self.exported()
        ws, r = s.find(f"{SRC}|UI/Mods/A")
        values = [c.value for c in ws[r]]
        target = s.wb["その他"]
        target.append(values)  # the same row pasted into another sheet, then edited there
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "二つ目で直した", nth=1)
        s.save()
        out = self.do_import()
        self.assertIn(f"PROBLEM {SRC}|UI/Mods/A: ID が重複しているので、「その他」の 5 行目は取り込みませんでした", out.out)
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("MOD A訳"))

    def test_duplicate_ids_with_differing_hidden_values(self):
        s = self.exported()
        ws, r = s.find(f"{SRC}|UI/Mods/A")
        values = [c.value for c in ws[r]]
        values[9] = "別の書き出し時の値"
        s.wb["その他"].append(values)
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "一つ目で直した")
        s.save()
        out = self.do_import()
        self.assertIn(f"PROBLEM {SRC}|UI/Mods/A: 書き出し時の値（非表示の列）が無いか食い違っているので取り込みませんでした", out.out)
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("MOD A訳"))

    def test_sorting_visible_columns_is_safe(self):
        s = self.exported()
        ws = s.wb["MOD画面"]
        # sort A-H only (Excel's sort of the filtered range): reverse the visible part of rows 2-5
        visible = [[ws.cell(r, c).value for c in range(1, 9)] for r in range(2, 6)]
        for i, vals in enumerate(reversed(visible)):
            for c, v in enumerate(vals, 1):
                ws.cell(2 + i, c).value = v
        s.set(f"{SRC}|UI/Mods/B", "日本語訳", "並べ替え後に直した")
        s.save()
        r = self.do_import()
        self.assertIn("修正 1 件", self.stats(r))
        self.assertEqual(self.json()["UI/Mods/B"]["ja"], "並べ替え後に直した")
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("MOD A訳"))

    def test_moving_columns_is_safe(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "列を動かしたシートで直した")
        ws = s.wb["MOD画面"]
        ws.move_range("B1:C10", cols=20)  # ファイル/種類 moved far to the right (ID stays in column A)
        s.save()
        self.do_import()
        self.assertEqual(self.json()["UI/Mods/A"]["ja"], "列を動かしたシートで直した")

    def test_column_inserted_before_id_is_imported(self):
        """A reviewer inserts a column of their own at A: the sheet is still read (columns by header name, RS3 fixed)."""
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "列を挿入したシートで直した")
        ws = s.wb["MOD画面"]
        ws.insert_cols(1)
        ws.cell(1, 1).value = "自分用メモ"
        ws.cell(3, 1).value = "レビュアー自身のメモ（取り込まれない）"
        s.save()
        r = self.do_import()
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("列を挿入したシートで直した", "edited"))
        self.assertEqual(self.stats(r), f"修正 1 件、確認済み 0 件、メモ 0 件、変更なし {TRANSLATED - 1} 件、競合 0 件")
        # the rows of that sheet are all read: the untranslated row is reported as before
        self.assertEqual(r.lines("PROBLEM"), [f"PROBLEM {SRC}|UI/Mods/D: 翻訳データにありません"])
        self.assertEqual(r.lines("WARNING"), [])
        self.assertEqual(self.json()["UI/Mods/B"], ours()[SRC]["UI/Mods/B"])

    def test_columns_inserted_between_and_after_the_fields(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/B", "状態", "確認済み")
        s.set(f"{SRC}|UI/Mods/C", "日本語訳", "途中に列を足したシートで直した")
        ws = s.wb["MOD画面"]
        ws.insert_cols(6)  # between 日本語訳 and 状態
        ws.cell(1, 6).value = "候補"
        ws.insert_cols(10)  # before the hidden columns
        ws.cell(1, 10).value = None  # a column without a header
        ws.cell(2, 10).value = "見出しの無い列"
        s.save()
        r = self.do_import()
        self.assertEqual(self.stats(r), f"修正 1 件、確認済み 1 件、メモ 0 件、変更なし {TRANSLATED - 2} 件、競合 0 件")
        got = self.json()
        self.assertEqual(got["UI/Mods/B"]["status"], "reviewed")
        self.assertEqual(got["UI/Mods/C"]["ja"], "途中に列を足したシートで直した")
        self.assertEqual(r.lines("WARNING"), [])

    def test_sheet_without_headers_is_reported(self):
        """A review sheet whose ID header was renamed is not imported, and a WARNING names it."""
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "見出しを消したシートで直した")
        s.set(f"{SRC}|X/3", "日本語訳", "その他シートで直した")
        s.wb["MOD画面"].cell(1, 1).value = "番号"
        s.save()
        r = self.do_import()
        self.assertIn("WARNING: シート「MOD画面」は ID・日本語訳・状態・訳注・レビューメモの見出しが無いので取り込みませんでした",
                      r.out)
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("MOD A訳"))
        self.assertEqual(self.json()["X/3"], h.entry("その他シートで直した", "edited"))  # the other sheets are read
        self.assertNotIn("使い方", "\n".join(r.lines("WARNING")))

    def test_sheet_missing_a_field_column_is_reported_not_crashing(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "列を消したシートで直した")
        s.wb["MOD画面"].delete_cols(8)  # レビューメモ
        s.save()
        r = self.do_import()
        self.assertIn("WARNING: シート「MOD画面」は ID・日本語訳・状態・訳注・レビューメモの見出しが無いので取り込みませんでした",
                      r.out)
        self.assertNotIn("Traceback", r.out)
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("MOD A訳"))

    def test_extra_worksheet_of_the_reviewer_is_reported(self):
        s = self.exported()
        extra = s.wb.create_sheet("自分用")
        extra.append(["気になった点"])
        extra.append(["UI/Mods/A の訳"])
        s.set(f"{SRC}|UI/Mods/B", "日本語訳", "直した")
        s.save()
        r = self.do_import()
        self.assertEqual(r.lines("WARNING"),
                         ["WARNING: シート「自分用」は ID・日本語訳・状態・訳注・レビューメモの見出しが無いので取り込みませんでした"])
        self.assertIn("修正 1 件", self.stats(r))

    def test_unknown_id(self):
        s = self.exported()
        s.wb["その他"].append(["I2LS_Test|Nope", "", "", "", "x", "未確認", "", "", "I2LS_Test|Nope", "x", "未確認", "", ""])
        s.save()
        r = self.do_import()
        self.assertIn("PROBLEM I2LS_Test|Nope: 翻訳データにありません", r.out)


def rewrite_zip(path, old: bytes, new: bytes):
    """Replace bytes inside the worksheet XML of an .xlsx (to produce what Excel itself would write)."""
    import zipfile
    tmp = path.with_name(path.name + ".rewrite")
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                data = data.replace(old, new)
            zout.writestr(info, data)
    os.replace(tmp, path)


def to_shared_strings(path, marker: str):
    """Move every inline-string cell containing `marker` into xl/sharedStrings.xml, as Excel saves text cells.

    The marker is removed from the text; the rest of the cell text (escapes included) is kept verbatim.
    """
    import re
    import zipfile
    tmp = path.with_name(path.name + ".sst")
    shared = []

    def move(m):
        attrs = (m.group(1) + m.group(2)).strip()
        shared.append(m.group(4).replace(marker, ""))
        return f'<c {attrs} t="s"><v>{len(shared) - 1}</v></c>'

    cell_re = re.compile(r'<c ([^>]*?)t="inlineStr"([^>]*)><is><t([^>]*)>([^<]*' + re.escape(marker) + r'[^<]*)</t></is></c>')
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        names = zin.namelist()
        assert "xl/sharedStrings.xml" not in names
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename.startswith("xl/worksheets/"):
                data = cell_re.sub(move, data.decode("utf-8")).encode("utf-8")
            elif info.filename == "[Content_Types].xml":
                data = data.replace(b"</Types>", b'<Override PartName="/xl/sharedStrings.xml" ContentType="application/'
                                    b'vnd.openxmlformats-officedocument.spreadsheetml.sharedStrings+xml"/></Types>')
            zout.writestr(info, data)
        sst = ('<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
               '<sst xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" '
               f'count="{len(shared)}" uniqueCount="{len(shared)}">'
               + "".join(f'<si><t xml:space="preserve">{t}</t></si>' for t in shared) + "</sst>")
        zout.writestr("xl/sharedStrings.xml", sst.encode("utf-8"))
    assert shared, "marker not found"
    os.replace(tmp, path)


class ExcelCarriageReturn(ReviewBase):
    """Excel stores a CR inside a cell as the escape _x000D_; openpyxl returns it as that literal text (RS2 fixed)."""

    def read_cell(self, sheet, ref):
        wb = load_workbook(self.xlsx, read_only=True)  # read-only mode keeps the file open until close()
        try:
            return wb[sheet][ref].value
        finally:
            wb.close()

    def test_crlf_typed_in_excel_becomes_a_line_break(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "直した@@CR@@\n改行")
        s.save()
        rewrite_zip(self.xlsx, b"@@CR@@", b"_x000D_")
        self.assertEqual(self.read_cell("MOD画面", "E2"), "直した_x000D_\n改行")
        r = self.do_import()
        self.assertEqual(self.json()["UI/Mods/A"], h.entry("直した\n改行", "edited"))
        self.assertIn("修正 1 件", self.stats(r))

    def test_crlf_in_shared_strings_as_excel_saves_it(self):
        """Excel keeps cell text in xl/sharedStrings.xml; a CR there is _x000D_ too."""
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "@@SST@@共有_x000D_\n文字列")
        s.set(f"{SRC}|UI/Mods/B", "訳注", "@@SST@@訳注_x000D_\n二行目")
        s.set(f"{SRC}|X/3", "レビューメモ", "@@SST@@メモ_x000D_\n二行目")
        s.save()
        to_shared_strings(self.xlsx, "@@SST@@")
        self.assertEqual(self.read_cell("MOD画面", "E2"), "共有_x000D_\n文字列")
        r = self.do_import()
        got = self.json()
        self.assertEqual(got["UI/Mods/A"], h.entry("共有\n文字列", "edited"))
        self.assertEqual(got["UI/Mods/B"], h.entry("MOD B訳", note="訳注\n二行目"))
        self.assertEqual(got["X/3"], h.entry("三つ目", review="メモ\n二行目"))
        self.assertNotIn("_x000D_", json.dumps(h.load_ours(self.root), ensure_ascii=False))
        self.assertIn("修正 1 件", self.stats(r))

    def test_cr_in_note_and_memo_and_tab_escape(self):
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/B", "訳注", "訳注@@CR@@\n続き")
        s.set(f"{SRC}|X/3", "レビューメモ", "タブ@@TAB@@区切り")
        s.save()
        rewrite_zip(self.xlsx, b"@@CR@@", b"_x000D_")
        rewrite_zip(self.xlsx, b"@@TAB@@", b"_x0009_")
        self.do_import()
        got = self.json()
        self.assertEqual(got["UI/Mods/B"]["note"], "訳注\n続き")
        self.assertEqual(got["X/3"]["review"], "タブ\t区切り")

    def test_escaped_underscore_stays_literal(self):
        """Excel writes a typed "_x000D_" as "_x005F_x000D_": that is text, not a line break."""
        s = self.exported()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "文字@@US@@x000D_のまま")
        s.save()
        rewrite_zip(self.xlsx, b"@@US@@", b"_x005F_")
        self.do_import()
        self.assertEqual(self.json()["UI/Mods/A"]["ja"], "文字_x000D_のまま")

    def test_escapes_in_untouched_rows_are_not_edits(self):
        """The hidden export-time columns are unescaped the same way, so a CR saved there is not a change."""
        s = self.exported()
        for name in ("日本語訳", "書き出し時の日本語訳"):
            s.set(f"{SRC}|X/1", name, "  一行目@@CR@@\n二行目\n")
        s.save()
        rewrite_zip(self.xlsx, b"@@CR@@", b"_x000D_")
        before = h.snapshot_files(self.root / "translations")
        r = self.do_import()
        self.assertEqual(self.stats(r), f"修正 0 件、確認済み 0 件、メモ 0 件、変更なし {TRANSLATED} 件、競合 0 件")
        self.assertEqual(h.snapshot_files(self.root / "translations"), before)


class CellText(unittest.TestCase):
    """review_sheet.cell_text() / snapshot() on fake cells (in-process)."""

    @classmethod
    def setUpClass(cls):
        cls.rs = h.load_tool("review_sheet")

    def cell(self, value, data_type="s"):
        return type("Cell", (), {"value": value, "data_type": data_type})()

    def test_table(self):
        ct = self.rs.cell_text
        self.assertEqual(ct("日本語訳", self.cell(None)), "")
        self.assertEqual(ct("日本語訳", self.cell("a\r\nb\rc")), "a\nb\nc")
        self.assertEqual(ct("日本語訳", self.cell("  前後  ")), "  前後  ")
        self.assertEqual(ct("状態", self.cell(" 確認済み\n")), "確認済み")
        self.assertIsNone(ct("日本語訳", self.cell("=A1", "f")))
        self.assertIsNone(ct("日本語訳", self.cell("=A1")))
        self.assertIsNone(ct("日本語訳", self.cell("#N/A", "e")))
        self.assertIsNone(ct("日本語訳", self.cell(12, "n")))
        self.assertIsNone(ct("日本語訳", self.cell(True, "b")))
        self.assertIsNone(ct("日本語訳", self.cell(datetime.datetime(2026, 1, 1), "d")))

    def test_snapshot(self):
        self.assertEqual(self.rs.snapshot({}), {"日本語訳": "", "状態": "未確認", "訳注": "", "レビューメモ": ""})
        self.assertEqual(self.rs.snapshot({"ja": "a", "status": "bogus"})["状態"], "未確認")
        self.assertEqual(self.rs.snapshot(h.entry("a", "edited", "n", review="m")),
                         {"日本語訳": "a", "状態": "修正済み", "訳注": "n", "レビューメモ": "m"})


class Legacy(ReviewBase):
    def legacy_sheet(self, created: str | None = None):
        s = self.exported()
        for ws in s.wb.worksheets:
            if ws.cell(1, 1).value == "ID":
                ws.delete_cols(9, 5)
        if created:
            for c in s.wb["使い方"]["A"]:
                if isinstance(c.value, str) and c.value.startswith("作成日時: "):
                    c.value = "作成日時: " + created
        return s

    def test_old_format_needs_legacy(self):
        s = self.legacy_sheet()
        s.save()
        before = h.snapshot_files(self.root / "translations")
        r = h.run_tool("review_sheet.py", "import", self.xlsx, project=self.root)
        self.assertEqual(r.code, 1, r)
        self.assertIn("--legacy", r.out)
        self.assertIn("MOD画面", r.out)
        self.assertEqual(h.snapshot_files(self.root / "translations"), before)

    def test_legacy_two_way_import(self):
        s = self.legacy_sheet()
        s.set(f"{SRC}|UI/Mods/A", "日本語訳", "レガシーで直した")
        s.set(f"{SRC}|UI/Mods/B", "状態", "確認済み")
        s.save()
        j = self.json()
        j["X/3"]["ja"] = "書き出し後にJSONで直した"
        self.save_json(SRC, j)
        r = self.do_import("--legacy")
        self.assertIn("WARNING: 古い形式のシート", r.out)
        self.assertIn("修正 2 件、確認済み 1 件", self.stats(r))
        got = self.json()
        self.assertEqual(got["UI/Mods/A"], h.entry("レガシーで直した", "edited"))
        self.assertEqual(got["UI/Mods/B"]["status"], "reviewed")
        # documented limit of --legacy: a JSON change after export goes back to the sheet value
        self.assertEqual(got["X/3"], h.entry("三つ目", "edited"))

    def test_legacy_warns_about_newer_translation_files(self):
        s = self.legacy_sheet(created="2020-01-01 00:00")
        s.save()
        r = self.do_import("--legacy")
        self.assertIn("WARNING: シートの作成日時（2020-01-01 00:00）より後に更新された翻訳データ: I2LS_Other.json, I2LS_Test.json",
                      r.out)

    def test_legacy_no_warning_for_older_files(self):
        s = self.legacy_sheet()
        s.save()
        old = time.time() - 3600
        for p in (self.root / "translations" / "ja").iterdir():
            os.utime(p, (old, old))
        r = self.do_import("--legacy")
        self.assertNotIn("より後に更新された翻訳データ", r.out)


@unittest.skipUnless(h.REAL_STRINGS.exists() and (h.PROJECT / "work" / "worklist.json").exists()
                     and (h.PROJECT / "translations" / "ja").exists(), "real project data missing")
class RealData(unittest.TestCase):
    def test_real_translations_no_op_round_trip(self):
        root = h.new_dir(self)
        (root / "work" / "extract").mkdir(parents=True)
        shutil.copyfile(h.REAL_STRINGS, root / "work" / "extract" / "strings.json")
        shutil.copyfile(h.PROJECT / "work" / "worklist.json", root / "work" / "worklist.json")
        shutil.copytree(h.PROJECT / "translations" / "ja", root / "translations" / "ja")
        xlsx = root / "review" / "real.xlsx"
        r = h.run_tool("review_sheet.py", "export", "--out", xlsx, project=root)
        self.assertEqual(r.code, 0, r)
        before = h.snapshot_files(root / "translations")
        r = h.run_tool("review_sheet.py", "import", xlsx, project=root)
        self.assertEqual(r.code, 0, r)
        self.assertEqual(r.lines("CONFLICT"), [])
        self.assertTrue(all("翻訳データにありません" in p for p in r.lines("PROBLEM")), r.lines("PROBLEM")[:5])
        self.assertRegex(r.out, r"修正 0 件、確認済み 0 件、メモ 0 件、変更なし \d+ 件、競合 0 件")
        self.assertEqual(h.snapshot_files(root / "translations"), before)
        worklist = {(w["source"], w["key"]) for w in json.loads((root / "work" / "worklist.json").read_text("utf-8"))}
        in_sheet = sum(1 for s, es in h.load_ours(root).items() for k in es if (s, k) in worklist)
        self.assertIn(f"変更なし {in_sheet} 件", r.out)


if __name__ == "__main__":
    unittest.main()
