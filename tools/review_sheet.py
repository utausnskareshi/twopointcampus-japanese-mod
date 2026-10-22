"""Review workbook for the Japanese drafts (Excel), and import of the reviewer's changes.

export: review/TPC_JP_review.xlsx — one sheet per category; columns ID, file, kind, English,
        Japanese, state (dropdown), translator note, reviewer memo; hidden columns I-M keep the exported
        ID/Japanese/state/note/memo. An existing file is not overwritten without --force.
import: three-way: only cells the reviewer changed since export are applied; if the JSON changed too, the row
        is reported as CONFLICT and left alone. Changed Japanese -> status "edited"; state 確認済み -> "reviewed";
        memo kept as "review". Blank/unknown states, formulas, numbers, dates and duplicate IDs are PROBLEMs.
        --legacy: also import sheets without the hidden columns (old two-way comparison with the current JSON).

usage: python tools/review_sheet.py export [--out FILE] [--force]
       python tools/review_sheet.py import FILE [--legacy]
"""
import datetime
import json
import sys
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils.escape import unescape
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from tpcloc import paths, translations  # noqa: E402

DEFAULT_OUT = paths.PROJECT_ROOT / "review" / "TPC_JP_review.xlsx"
SHEETS = [("Mods", "MOD画面"), ("DLC_Hospital", "メディカルスクール"), ("DLC_Ghost", "スクールスピリッツ"),
          ("ChallengeMode", "チャレンジモード"), ("Seasonal", "季節イベント"), ("Promo", "コラボ・特典"),
          ("Cafeteria", "カフェテリア"), ("Other", "その他")]
STATE_LABELS = {"draft": "未確認", "reviewed": "確認済み", "edited": "修正済み"}
LABEL_STATES = {v: k for k, v in STATE_LABELS.items()}
HEADERS = ["ID", "ファイル", "種類", "英語（原文）", "日本語訳", "状態", "訳注", "レビューメモ"]
WIDTHS = [22, 16, 9, 60, 60, 11, 34, 30]
FIELDS = ["日本語訳", "状態", "訳注", "レビューメモ"]
# hidden columns after HEADERS: the values at export time, looked up by their own ID (so sorting A-H is safe)
BASE_HEADERS = ["書き出し時ID"] + [f"書き出し時の{n}" for n in FIELDS]
KIND_LABELS = {"name": "名前", "title": "見出し", "label": "ラベル", "text": "文章", "person_name": "人名",
               "tannoy": "構内放送", "radio": "ラジオ"}

USAGE = [
    "Two Point Campus 日本語化パッチ　下訳レビュー用シート",
    "",
    "■ 見かた",
    "・カテゴリごとにシートを分けています（MOD画面、メディカルスクール、スクールスピリッツ、チャレンジモード など）。",
    "・「英語（原文）」と「日本語訳」のセル内改行は、ゲーム内の改行と同じ位置です。",
    "・{COUNT} や <style=\"BodyBold\">…</style>、<br>、[i2p_One] などの記号はゲームが使う印なので、消したり書き換えたりしないでください（位置の移動は可）。",
    "・「訳注」は下訳時のメモです（ダジャレの原意、英文の誤記、判断に迷った点など）。",
    "",
    "■ レビューのしかた",
    "・訳を直すときは「日本語訳」のセルを直接書き換えてください（取り込み時に自動で「修正済み」になります）。",
    "・そのままで良い行は「状態」を「確認済み」にしてください。",
    "・気になる点だけ書きたいときは「レビューメモ」に書いてください（訳は変わりません）。",
    "・行の並べ替えやフィルターは自由です。ID 列は消さないでください。",
    "・セルに数式（=E2 など）は使わず、文字で入力してください。非表示の列（書き出し時の値）は取り込みに使うので、変えないでください。",
    "",
    "■ 取り込み（作業者向け）",
    "python tools/review_sheet.py import review/TPC_JP_review.xlsx → python tools/qa.py → python tools/build_patch.py",
]


def opt(name, default=None):
    if name in sys.argv:
        return sys.argv[sys.argv.index(name) + 1]
    return default


def snapshot(e: dict) -> dict[str, str]:
    """The reviewable values of an entry as the sheet shows them (keys = FIELDS)."""
    return {"日本語訳": e.get("ja", ""), "状態": STATE_LABELS.get(e.get("status", "draft"), "未確認"),
            "訳注": e.get("note", ""), "レビューメモ": e.get("review", "")}


def cell_text(name: str, cell) -> str | None:
    """Cell value as text with \\n line breaks; None for a formula, an error value, a number or a date."""
    v = cell.value
    if v is None:
        return ""
    if cell.data_type in ("f", "e") or not isinstance(v, str) or v.startswith("="):
        return None
    # Excel stores some characters as OOXML escapes (_x000D_ for CR) that openpyxl returns literally
    v = unescape(v).replace("\r\n", "\n").replace("\r", "\n")
    return v.strip() if name == "状態" else v


def show(v: str) -> str:
    return v.replace("\n", "↵")


def export(out: Path, force: bool = False):
    if out.exists() and not force:
        sys.exit(f"{out} はすでにあります（レビュー中のシートかもしれません）。上書きするときは --force を付けてください")
    worklist = json.loads((paths.WORK_DIR / "worklist.json").read_text("utf-8"))
    ours = translations.load_all()
    wb = Workbook()
    ws = wb.active
    ws.title = "使い方"
    for i, line in enumerate(USAGE + ["", f"作成日時: {datetime.datetime.now():%Y-%m-%d %H:%M}"], 1):
        ws.cell(row=i, column=1, value=line).font = Font(bold=(i == 1), size=12 if i == 1 else 11)
    ws.column_dimensions["A"].width = 120

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    id_font = Font(color="808080", size=9)
    wrap = Alignment(wrap_text=True, vertical="top")
    counts = {}
    for category, title in SHEETS:
        rows = [w for w in worklist if w["category"] == category]
        if not rows:
            continue
        sh = wb.create_sheet(title)
        sh.append(HEADERS + BASE_HEADERS)
        for col, width in enumerate(WIDTHS, 1):
            sh.column_dimensions[sh.cell(row=1, column=col).column_letter].width = width
            c = sh.cell(row=1, column=col)
            c.font = Font(bold=True)
            c.fill = header_fill
        for col in range(len(HEADERS) + 1, len(HEADERS) + len(BASE_HEADERS) + 1):
            sh.column_dimensions[sh.cell(row=1, column=col).column_letter].hidden = True
        dv = DataValidation(type="list", formula1='"未確認,確認済み,修正済み"', allow_blank=False,
                            showErrorMessage=True, error="未確認・確認済み・修正済みから選んでください")
        sh.add_data_validation(dv)
        for w in rows:
            e = ours.get(w["source"], {}).get(w["key"], {})
            ref, snap = f"{w['source']}|{w['key']}", snapshot(e)
            sh.append([ref, w["source"].replace("I2LS_", ""), KIND_LABELS.get(w["kind"], w["kind"]), w["en"]]
                      + [snap[n] for n in FIELDS] + [ref] + [snap[n] for n in FIELDS])
            r = sh.max_row
            for col in range(1, 9):
                sh.cell(row=r, column=col).alignment = wrap
            sh.cell(row=r, column=1).font = id_font
            dv.add(sh.cell(row=r, column=6))
        sh.freeze_panes = "D2"
        sh.auto_filter.ref = f"A1:H{sh.max_row}"
        counts[title] = len(rows)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    print(f"{sum(counts.values())} rows -> {out}")
    print(counts)


def sheet_created(wb) -> datetime.datetime | None:
    """The 作成日時 line of the 使い方 sheet (minute precision)."""
    if "使い方" not in wb.sheetnames:
        return None
    for row in wb["使い方"].iter_rows(max_col=1, values_only=True):
        if isinstance(row[0], str) and row[0].startswith("作成日時: "):
            try:
                return datetime.datetime.strptime(row[0][len("作成日時: "):], "%Y-%m-%d %H:%M")
            except ValueError:
                return None
    return None


def read_baseline(sheets) -> dict:
    """Exported values (hidden columns) by ID; None for an ID whose copies differ or cannot be read."""
    base = {}
    for sh, col, width in sheets:
        if not all(h in col for h in BASE_HEADERS):
            continue
        for cells in sh.iter_rows(min_row=2, max_col=width):
            ref = cells[col[BASE_HEADERS[0]]].value
            if not ref:
                continue
            vals = {n: cell_text(n, cells[col[f"書き出し時の{n}"]]) for n in FIELDS}
            if None in vals.values() or base.get(ref, vals) != vals:
                vals = None
            base[ref] = vals
    return base


def import_sheet(path: Path, legacy: bool = False):
    data = json.loads((paths.WORK_DIR / "extract" / "strings.json").read_text("utf-8"))
    ours = translations.load_all()
    # data_only stays False: a formula then reads as "=..." and is refused instead of importing its cached result
    wb = load_workbook(path, read_only=True)
    sheets = []
    for sh in wb.worksheets:
        header = [c.value for c in next(sh.iter_rows(max_row=1), ())]
        # columns are found by name, so a column inserted by the reviewer does not matter
        if "ID" in header and all(n in header for n in FIELDS):
            sheets.append((sh, {name: i for i, name in enumerate(header) if name is not None}, len(header)))
        elif sh.title != "使い方":
            print(f"WARNING: シート「{sh.title}」は ID・日本語訳・状態・訳注・レビューメモの見出しが無いので取り込みませんでした")
    old_format = [sh.title for sh, col, _ in sheets if not all(h in col for h in BASE_HEADERS)]
    if old_format and not legacy:
        sys.exit(f"{path} は書き出し時の値（非表示の列）が無い古い形式のシートです（{'、'.join(old_format)}）。"
                 "書き出し直したシートを使うか、--legacy を付けて取り込んでください")
    if old_format:
        print("WARNING: 古い形式のシートは、シートと現在の翻訳データを比べて取り込みます。"
              "書き出し後に直した訳・状態・訳注・メモは、シートの値に戻ります")
        created = sheet_created(wb)
        newer = [f.name for f in sorted(paths.TRANSLATIONS_DIR.glob("*.json"))
                 if created and f.stat().st_mtime >= created.timestamp() + 60]
        if newer:
            print(f"WARNING: シートの作成日時（{created:%Y-%m-%d %H:%M}）より後に更新された翻訳データ: {', '.join(newer)}")
    base = read_baseline(sheets)
    stats = {"edited": 0, "reviewed": 0, "memo": 0, "unchanged": 0, "conflict": 0}
    problems, conflicts, changed_sources, seen = [], [], set(), set()
    for sh, col, width in sheets:
        has_base = sh.title not in old_format
        for r, cells in enumerate(sh.iter_rows(min_row=2, max_col=width), 2):
            ref = cells[col["ID"]].value
            if not ref:
                continue
            source, _, key = str(ref).partition("|")
            entry = ours.get(source, {}).get(key)
            if entry is None:
                problems.append(f"{ref}: 翻訳データにありません")
                continue
            if ref in seen:
                problems.append(f"{ref}: ID が重複しているので、「{sh.title}」の {r} 行目は取り込みませんでした")
                continue
            seen.add(ref)
            cell = {n: cell_text(n, cells[col[n]]) for n in FIELDS}
            bad = [n for n in FIELDS if cell[n] is None]
            if bad:
                problems.append(f"{ref}: {'・'.join(bad)}が数式・数値・日付などになっているので取り込みませんでした"
                                "（文字で入力し直してください）")
                continue
            ja, label, note, memo = (cell[n] for n in FIELDS)
            if not ja.strip():
                problems.append(f"{ref}: 日本語訳が空なので取り込みませんでした")
                continue
            cur = snapshot(entry)
            b = base.get(ref) if has_base else cur  # --legacy: the current JSON stands in for the exported values
            if b is None:
                problems.append(f"{ref}: 書き出し時の値（非表示の列）が無いか食い違っているので取り込みませんでした")
                continue
            # three-way: a cell equal to b was not touched; a change is applied only if the JSON still equals b
            before, clash, counts = dict(entry), [], {"edited": 0, "reviewed": 0, "memo": 0}
            if ja != b["日本語訳"]:
                if cur["日本語訳"] == b["日本語訳"]:
                    entry["ja"] = ja
                    entry["status"] = "edited"
                    counts["edited"] += 1
                elif cur["日本語訳"] != ja:
                    clash.append(("日本語訳がシートと翻訳データの両方で変わっています", ja, b["日本語訳"], cur["日本語訳"]))
            elif label != b["状態"]:
                state = LABEL_STATES.get(label)
                if state is None:
                    problems.append(f"{ref}: 状態「{label}」は未確認・確認済み・修正済みのどれでもないので、状態は変えませんでした")
                elif cur["日本語訳"] != b["日本語訳"]:
                    clash.append((f"状態を「{label}」にした訳が、書き出し後に翻訳データで変わっています",
                                  ja, b["日本語訳"], cur["日本語訳"]))
                elif cur["状態"] not in (b["状態"], label):
                    clash.append(("状態がシートと翻訳データの両方で変わっています", label, b["状態"], cur["状態"]))
                elif state != entry.get("status"):
                    entry["status"] = state
                    if state == "reviewed":
                        counts["reviewed"] += 1
            if note != b["訳注"] and note != cur["訳注"]:
                if cur["訳注"] == b["訳注"]:
                    entry["note"] = note
                else:
                    clash.append(("訳注がシートと翻訳データの両方で変わっています", note, b["訳注"], cur["訳注"]))
            if memo != b["レビューメモ"] and memo != cur["レビューメモ"]:
                if cur["レビューメモ"] != b["レビューメモ"]:
                    clash.append(("レビューメモがシートと翻訳データの両方で変わっています",
                                  memo, b["レビューメモ"], cur["レビューメモ"]))
                elif memo:
                    entry["review"] = memo
                    counts["memo"] += 1
                else:
                    entry.pop("review", None)
            for what, new, old, now in clash:
                conflicts.append(f"{ref}: {what}（シート「{show(new)}」／書き出し時「{show(old)}」／現在「{show(now)}」）")
            if clash:
                # a row with a conflict is not imported at all (its other cells may depend on the conflicting one)
                entry.clear()
                entry.update(before)
                stats["conflict"] += 1
            else:
                for k, n in counts.items():
                    stats[k] += n
            if entry != before:
                changed_sources.add(source)
            elif not clash:
                stats["unchanged"] += 1
    for source in sorted(changed_sources):
        translations.save(source, ours[source], data["sources"][source]["keys"])
    print(f"修正 {stats['edited']} 件、確認済み {stats['reviewed']} 件、メモ {stats['memo']} 件、"
          f"変更なし {stats['unchanged']} 件、競合 {stats['conflict']} 件")
    for p in problems:
        print("PROBLEM", p)
    for c in conflicts:
        print("CONFLICT", c)
    if conflicts:
        print("CONFLICT の項目は取り込んでいません。シートと translations/ja を見比べて直してください")
    print("次に python tools/qa.py を実行してください")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    if sys.argv[1] == "export":
        export(Path(opt("--out", DEFAULT_OUT)), force="--force" in sys.argv)
    elif sys.argv[1] == "import" and len(sys.argv) > 2:
        import_sheet(Path(sys.argv[2]), legacy="--legacy" in sys.argv)
    else:
        sys.exit(__doc__)
